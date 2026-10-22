import sys,io,base64,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'engine'))
from core import Engine,b64,EngineError
import pymupdf as f
import pytest

def test_simple_object_preserves_other_streams():
 raw=Path('test-fixtures/research.pdf').read_bytes();e=Engine();doc=e.open(b64(raw));sid=doc['id'];s=e.objects(sid)['texts'][1]
 p=e.preview(sid,[{'type':'replaceText','id':s['id'],'text':s['text'].replace('should','must')}]);assert not p['warnings']
 e.commit(sid,p['token']);before=f.open(stream=raw,filetype='pdf');after=f.open(stream=e.session(sid).data,filetype='pdf')
 old=[before.xref_stream(x) for x in before[0].get_contents()];new=[after.xref_stream(x) for x in after[0].get_contents()]
 assert len(old)==len(new);assert sum(a!=b for a,b in zip(old,new))==1
 for i in [1,2]:assert before[i].get_pixmap().samples==after[i].get_pixmap().samples
 Path('tests/results/text-diff.json').write_text(json.dumps(p,indent=2))

def test_locales():
 a=json.loads(Path('src/i18n/en-US.json').read_text());b=json.loads(Path('src/i18n/zh-CN.json').read_text());assert a.keys()==b.keys();assert all(a.values()) and all(b.values())

def test_import_docx_and_image():
 from docx import Document
 from PIL import Image
 e=Engine();buf=io.BytesIO();d=Document();d.add_paragraph('Imported document');d.add_paragraph('中文段落');d.save(buf)
 s=e.import_file(b64(buf.getvalue()),'input.docx');assert 'Imported document' in ''.join(x['text'] for x in e.objects(s['id'])['texts'])
 b=io.BytesIO();Image.new('RGB',(80,60),'blue').save(b,format='PNG');s=e.import_file(b64(b.getvalue()),'image.png');assert len(e.objects(s['id'])['images'])==1

def test_reorder_add_delete_resize_and_search():
 e=Engine();s=e.open(b64(Path('test-fixtures/research.pdf').read_bytes()));sid=s['id']
 assert len(e.search(sid,'should'))==6
 for op in [{'type':'reorder','order':[2,0,1]},{'type':'resize','width':620,'height':900},{'type':'addPage'},{'type':'deletePage','page':1}]:
  p=e.preview(sid,[op]);e.commit(sid,p['token'])
 assert e.info(sid)['pages'][0]['width']==620
 assert 'Page 3' in ''.join(x['text'] for x in e.objects(sid)['texts'])

def test_header_watermark_addtext_and_underline():
 for op in [{'type':'pageNumbers'},{'type':'watermark','text':'DRAFT'},{'type':'headerFooter','text':'中文页眉'},{'type':'addText','text':'New text','rect':[50,400,300,450]},{'type':'underline','query':'should'}]:
  e=Engine();s=e.open(b64(Path('test-fixtures/research.pdf').read_bytes()));p=e.preview(s['id'],[op]);assert p['changes']>0;e.commit(s['id'],p['token'])

def test_overlap_edit_fails_closed():
 d=f.open();p=d.new_page();p.insert_text((60,100),'First line');p.insert_text((60,102),'Overlapping')
 e=Engine();s=e.open(b64(d.tobytes()));span=e.objects(s['id'])['texts'][0]
 # Fallback redaction may touch another overlapping line. Visual-diff guard must reject any resulting outside change.
 before=e.session(s['id']).data
 try:
  result=e.preview(s['id'],[{'type':'replaceText','id':span['id'],'text':'Short','color':'#cc0000'}])
  assert not any(x.get('outsideRegionChanged') for x in result['diff'])
 except EngineError as ex:assert ex.code in ['outsideRegionChanged','textOverflow','overlappingText']
 assert e.session(s['id']).data==before
