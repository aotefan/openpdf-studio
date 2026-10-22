"""Read actual UI-generated exports with independent document libraries."""
import json,zipfile,io
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image
root=Path('tests/results');manifest=json.loads((root/'batch-conversion-files-zh-CN.json').read_text())
assert set(manifest)=={'txt','docx','xlsx','pptx','png'}
report=[]
for kind,directory in manifest.items():
 files=sorted(Path(directory).iterdir());assert len(files)==2
 for i,p in enumerate(files):
  expected='OpenPDF' if i==0 else 'PAGE 1';pages=3 if i==0 else 6
  if kind=='txt':assert expected in p.read_text()
  elif kind=='docx':assert expected in '\n'.join(x.text for x in Document(p).paragraphs)
  elif kind=='xlsx':
   w=load_workbook(p);assert len(w.worksheets)==pages
   assert any(expected in str(c.value) for row in w.worksheets[0] for c in row)
  elif kind=='pptx':
   slides=Presentation(p).slides;assert len(slides)==pages
   assert all(any(s.shape_type==13 for s in slide.shapes) for slide in slides)
  else:
   with zipfile.ZipFile(p) as z:
    assert len(z.namelist())==pages
    for name in z.namelist():
     with Image.open(io.BytesIO(z.read(name))) as im:im.verify()
  report.append({'format':kind,'file':str(p),'valid':True})
(root/'batch-conversion-readback.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(f'{len(report)} actual exports reopened and verified across 5 formats.')
