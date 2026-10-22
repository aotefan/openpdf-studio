"""Independent local PDF operations. Candidate copies are committed only after preview."""
from __future__ import annotations
import base64, io, json, re, uuid, zipfile, subprocess, tempfile, os, math
from pathlib import Path
from dataclasses import dataclass, field
import pymupdf as fitz
from PIL import Image, ImageChops, ImageDraw

MAX_PAGES = 500

class EngineError(Exception):
    def __init__(self, code, details=None):
        self.code, self.details = code, details
        super().__init__(code)

def b64(data): return base64.b64encode(data).decode()
def un64(data):
    if not isinstance(data,str) or len(data)>100*1024*1024: raise EngineError('fileTooLarge')
    return base64.b64decode(data, validate=True)
def checked(data):
    try: d=fitz.open(stream=data,filetype='pdf')
    except Exception: raise EngineError('invalidPdf')
    if d.needs_pass or d.is_encrypted: d.close(); raise EngineError('encrypted')
    if not d.permissions & fitz.PDF_PERM_MODIFY: d.close(); raise EngineError('restricted')
    if not 0 < len(d) <= MAX_PAGES: d.close(); raise EngineError('pageLimit')
    if d.get_sigflags()>0: d.close(); raise EngineError('signedPdf')
    return d

def spans(page):
    result=[]
    for bi,b in enumerate(page.get_text('dict')['blocks']):
        for li,l in enumerate(b.get('lines',[])):
            for si,s in enumerate(l['spans']):
                result.append({'id':f'{bi}:{li}:{si}','text':s['text'],'rect':list(s['bbox']), 'viewRect':list(fitz.Rect(s['bbox'])*page.rotation_matrix),'origin':list(s['origin']),'size':s['size'],'font':s['font'],'color':'#%06x'%s['color'],'direction':list(l['dir'])})
    return result

def color(hexstr):
    if not re.fullmatch(r'#[0-9a-fA-F]{6}',hexstr): raise EngineError('invalidOperation')
    return tuple(int(hexstr[i:i+2],16)/255 for i in (1,3,5))
MARKUP_TYPES={'highlight':fitz.PDF_ANNOT_HIGHLIGHT,'underline':fitz.PDF_ANNOT_UNDERLINE,'strikeout':fitz.PDF_ANNOT_STRIKE_OUT}
MARKUP_TAG='keyword-markup-v1'
def validate_markup(query,style):
    if not isinstance(query,str) or not query.strip() or len(query)>200 or style not in MARKUP_TYPES:raise EngineError('invalidMarkup')
def quad_key(points):return tuple(round(v,2) for point in points for v in point)
def markup_hits(page,query,style):
    validate_markup(query,style);existing=set()
    for annot in page.annots() or []:
        if annot.type[0]!=MARKUP_TYPES[style]:continue
        vertices=annot.vertices or []
        for i in range(0,len(vertices)-3,4):existing.add(quad_key(vertices[i:i+4]))
    hits=[];duplicates=0;seen=set()
    for quad in page.search_for(query,quads=True):
        key=quad_key(quad)
        if key in seen:continue
        seen.add(key)
        if key in existing:duplicates+=1
        else:hits.append(quad)
    return hits,duplicates

def font_for(text, original='', bold=False, italic=False):
    if any(ord(c)>255 for c in text):return 'china-s'
    if bold and italic:return 'hebi'
    if bold:return 'hebo'
    if italic:return 'heit'
    if 'Times' in original:return 'tiro'
    if 'Courier' in original:return 'cour'
    return 'helv'
def replacement_layout(page,span,text,size,font,validate=True):
    """Grow from the original baseline. Redaction still uses ONLY the old glyphs."""
    if not math.isfinite(size) or not 3<=size<=144 or '\n' in text or '\r' in text:raise EngineError('invalidTextLayout')
    rect=fitz.Rect(span['rect']);origin=fitz.Point(span['origin'])
    metric=fitz.Font('cjk' if font=='china-s' else font)
    width=metric.text_length(text,fontsize=size)
    grown=fitz.Rect(rect)
    if text:
        grown.include_rect(fitz.Rect(origin.x,origin.y-size*metric.ascender,origin.x+width,origin.y-size*metric.descender))
    if validate and text:
        bounds=page.rect*page.derotation_matrix
        if not fitz.Rect(bounds.x0-.5,bounds.y0-.5,bounds.x1+.5,bounds.y1+.5).contains(grown):
            raise EngineError('textOverflow',{'required':width,'available':bounds.x1-origin.x})
        # Reject only newly occupied space. Never redact the expanded rectangle.
        for other in spans(page):
            if other['id']==span['id']:continue
            box=fitz.Rect(other['rect'])
            if (grown & box).get_area()>(rect & box).get_area()+.5:raise EngineError('textExpansionConflict')
        for other in page.get_image_info():
            box=fitz.Rect(other['bbox'])
            if not box.intersects(rect) and box.intersects(grown):raise EngineError('textExpansionConflict')
        # A vertical rule is a cell / layout boundary, not spare horizontal space.
        if grown.x1>rect.x1+.5:
            for path in page.get_drawings():
                edges=[]
                for item in path['items']:
                    if item[0]=='l':edges.append((item[1],item[2]))
                    elif item[0]=='re':
                        box=item[1];edges.extend([(box.tl,box.bl),(box.tr,box.br)])
                for a,b in edges:
                    if abs(a.x-b.x)<.5 and rect.x1+.5<a.x<grown.x1 and min(a.y,b.y)<grown.y1 and max(a.y,b.y)>grown.y0:
                        raise EngineError('textExpansionConflict')
    return {'rect':list(grown),'viewRect':list(grown*page.rotation_matrix),'width':width}

def embed(page,font):
    if font=='china-s':
        page.insert_font(fontname='OpenPDFCJK',fontbuffer=fitz.Font('cjk').buffer)
        return 'OpenPDFCJK'
    return font

def replace_simple_object(doc,page,span,text,op):
    """Conservative fast path: one isolated WinAnsi, standard-font text stream.

    No decoding assumptions for embedded fonts, forms, kerning arrays or shared streams.
    All graphics state and other page content streams remain byte-for-byte unchanged.
    """
    if any(op.get(k) for k in ['bold','italic','underline']):return False
    if abs(float(op.get('size',span['size']))-span['size'])>.001 or op.get('color',span['color'])!=span['color']:return False
    if not text or any(ord(c)<32 or ord(c)>126 for c in text+span['text']):return False
    candidates=[]
    for xref in page.get_contents():
        stream=doc.xref_stream(xref)
        if stream.count(b'BT')!=1 or stream.count(b'ET')!=1:continue
        shows=re.findall(rb'\[\s*<([0-9a-fA-F]+)>\s*\]\s*TJ',stream)
        fonts=re.findall(rb'/([A-Za-z0-9]+)\s+[\d.]+\s+Tf',stream)
        if len(shows)!=1 or len(fonts)!=1:continue
        if bytes.fromhex(shows[0].decode())!=span['text'].encode('ascii'):continue
        font=next((f for f in page.get_fonts() if f[4]==fonts[0].decode()),None)
        if not font or font[1]!='n/a' or font[2]!='Type1' or font[3]!=span['font'] or font[5]!='WinAnsiEncoding':continue
        if font[3] not in ['Helvetica','Helvetica-Bold','Helvetica-Oblique','Helvetica-BoldOblique','Times-Roman','Courier']:continue
        if sum(xref in p.get_contents() for p in doc)!=1:continue
        candidates.append((xref,stream,shows[0]))
    if len(candidates)!=1:return False
    xref,stream,old=candidates[0]
    doc.update_stream(xref,stream.replace(b'<'+old+b'>',b'<'+text.encode('ascii').hex().encode()+b'>',1))
    return True

def unique_image(doc,page,xref):
    occurrences=[(i,o) for i,p in enumerate(doc) for o in p.get_image_info(xrefs=True) if o['xref']==xref]
    if xref<=0 or len(occurrences)!=1 or occurrences[0][0]!=page.number:raise EngineError('sharedImage')
    return occurrences[0][1]

def image_placement(doc,page,xref):
    """Bounded direct image placement: edit its matrix in place, never reorder Do.

    Forms, clipping, rotated pages and non-isolated drawing streams are rejected.
    No mutation takes place until every placement and ownership check succeeds.
    """
    info=unique_image(doc,page,xref)
    if page.rotation or any(x['type']=='clip' for x in page.get_drawings(extended=True)):
        raise EngineError('imageGeometryUnsupported')
    transforms=info['transform']
    if abs(transforms[1])+abs(transforms[2])>.001 or transforms[0]<=0 or transforms[3]<=0:
        raise EngineError('imageGeometryUnsupported')
    names=[x[7].encode() for x in page.get_images(full=True) if x[0]==xref and x[9]==0]
    number=rb'[-+]?(?:\d*\.\d+|\d+\.?\d*)'
    pattern=rb'\s*q\s+('+rb'\s+'.join([number]*6)+rb')\s+cm\s+/([^\s/]+)\s+Do\s+Q\s*'
    matches=[]
    for content in page.get_contents():
        data=doc.xref_stream(content);match=re.fullmatch(pattern,data)
        if match and match[2] in names:
            matrix=fitz.Matrix(*map(float,match[1].split()))
            bbox=fitz.Rect(0,0,1,1)*matrix*page.transformation_matrix
            if any(abs(a-b)>.02 for a,b in zip(bbox,info['bbox'])):raise EngineError('imageGeometryUnsupported')
            if sum(content in p.get_contents() for p in doc)!=1:raise EngineError('sharedImage')
            matches.append((content,data,match,matrix))
    if len(matches)!=1:raise EngineError('imageGeometryUnsupported')
    return matches[0]

def move_image(doc,page,xref,old,new):
    content,data,match,matrix=image_placement(doc,page,xref)
    sx,sy=new.width/old.width,new.height/old.height
    change=fitz.Matrix(sx,0,0,sy,new.x0-old.x0*sx,new.y0-old.y0*sy)
    matrix=matrix*page.transformation_matrix*change*~page.transformation_matrix
    encoded=' '.join(format(v,'.8f').rstrip('0').rstrip('.') or '0' for v in matrix).encode()
    doc.update_stream(content,data[:match.start(1)]+encoded+data[match.end(1):])

def image_png(doc,xref):
    extracted=doc.extract_image(xref)
    if extracted.get('smask'):
        return fitz.Pixmap(fitz.Pixmap(doc,xref),fitz.Pixmap(doc,extracted['smask'])).tobytes('png')
    with Image.open(io.BytesIO(extracted['image'])) as im:
        buf=io.BytesIO();im.convert('RGBA').save(buf,format='PNG');return buf.getvalue()

def png(page,scale=1):return page.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False).tobytes('png')

def diff_report(before,after,regions,structural=False,alpha=False):
    with fitz.open(stream=before,filetype='pdf') as a,fitz.open(stream=after,filetype='pdf') as b:
        rows=[]
        for i in range(max(len(a),len(b))):
            if i>=len(a) or i>=len(b) or a[i].rect != b[i].rect:
                rows.append({'page':i,'changed':True,'structural':True});continue
            if alpha:
                pa=a[i].get_pixmap(alpha=True);pb=b[i].get_pixmap(alpha=True)
                aa=Image.frombytes('RGBA',(pa.width,pa.height),pa.samples);bb=Image.frombytes('RGBA',(pb.width,pb.height),pb.samples)
            else:aa=Image.open(io.BytesIO(png(a[i]))).convert('RGB');bb=Image.open(io.BytesIO(png(b[i]))).convert('RGB')
            delta=ImageChops.difference(aa,bb);combined=delta.getchannel(0)
            for channel in delta.split()[1:]:combined=ImageChops.lighter(combined,channel)
            box=combined.getbbox()
            if box:
                outside=delta.copy(); draw=ImageDraw.Draw(outside)
                for r in regions.get(i,[]):
                    rr=fitz.Rect(r)*a[i].rotation_matrix
                    draw.rectangle((max(0,rr.x0-3),max(0,rr.y0-3),rr.x1+3,rr.y1+3),fill=(0,)*len(delta.getbands()))
                outside_max=max(v[1] for v in outside.getextrema())
                unexpected=outside_max>1 and not structural
                rows.append({'page':i,'changed':True,'bounds':list(box),'structural':structural,'outsideRegionChanged':unexpected,'outsideMaxChannelDelta':outside_max,'outsideExact':outside_max==0})
        return rows

def recognize_page(page):
    """Local Vision result becomes an invisible text layer; never alter the scan."""
    if page.rotation:raise EngineError('ocrRotation')
    if page.get_text().strip():raise EngineError('alreadyText')
    helper=Path(os.environ.get('OPENPDF_OCR',str(Path(__file__).parent.parent/'build/vision-ocr')))
    if not helper.exists():raise EngineError('ocrUnavailable')
    try:
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'page.png';path.write_bytes(png(page,2))
            result=subprocess.run([str(helper),str(path)],capture_output=True,text=True,timeout=120)
        if result.returncode:raise EngineError('ocrFailed')
        lines=json.loads(result.stdout)
        if not isinstance(lines,list):raise ValueError()
        texts=[];low=0
        for item in lines:
            text=item['text'];x,y,w,h=item['rect'];confidence=float(item.get('confidence',1))
            if not isinstance(text,str) or any(not math.isfinite(v) for v in [x,y,w,h,confidence]) or not (0<=x<=1 and 0<=y<=1 and 0<w<=1 and 0<h<=1 and x+w<=1.001 and y+h<=1.001):raise ValueError()
            if not text.strip():continue
            r=fitz.Rect(x*page.rect.width,(1-y-h)*page.rect.height,(x+w)*page.rect.width,(1-y)*page.rect.height)
            size=max(1,r.height*.8);font=font_for(text);metric=fitz.Font('cjk' if font=='china-s' else font)
            width=metric.text_length(text,fontsize=size)
            if width>r.width:size*=r.width/width
            page.insert_text((r.x0,r.y1-r.height*.15),text,fontname=embed(page,font),fontsize=size,render_mode=3)
            texts.append(text);low+=confidence<.8
        return {'lines':len(texts),'lowConfidence':low,'sample':'\n'.join(texts[:3])[:360]}
    except subprocess.TimeoutExpired:raise EngineError('ocrTimeout')
    except EngineError:raise
    except (ValueError,KeyError,TypeError,OSError):raise EngineError('ocrFailed')

@dataclass
class Session:
    name:str
    data:bytes
    revision:int=0
    undo:list=field(default_factory=list)
    redo:list=field(default_factory=list)
    log:list=field(default_factory=list)

class Engine:
    def __init__(self):self.sessions={};self.previews={};self.ocr_jobs={}
    def open(self,data,name='Document.pdf'):
        raw=un64(data)
        with checked(raw):pass
        if len(self.sessions)>=12:raise EngineError('sessionLimit')
        sid=uuid.uuid4().hex;self.sessions[sid]=Session(Path(name).name,raw)
        return self.info(sid)
    def session(self,sid):
        if sid not in self.sessions:raise EngineError('sessionMissing')
        return self.sessions[sid]
    def info(self,sid):
        s=self.session(sid)
        with fitz.open(stream=s.data,filetype='pdf') as d:
            return {'id':sid,'name':s.name,'revision':s.revision,'pages':[{'width':p.rect.width,'height':p.rect.height,'rotation':p.rotation} for p in d],'bookmarks':d.get_toc(),'canUndo':bool(s.undo),'canRedo':bool(s.redo),'log':s.log[-30:]}
    def close(self,sid):
        self.sessions.pop(sid,None)
        self.ocr_jobs={k:v for k,v in self.ocr_jobs.items() if v['sid']!=sid}
        self.previews={k:v for k,v in self.previews.items() if v['sid']!=sid}
        return True
    def ocr_plan(self,sid,pages):
        s=self.session(sid)
        with fitz.open(stream=s.data,filetype='pdf') as d:
            if not isinstance(pages,list) or not pages or any(type(i) is not int or not 0<=i<len(d) for i in pages):raise EngineError('invalidPage')
            rows=[]
            for i in sorted(set(pages)):
                p=d[i]
                status='skipText' if p.get_text().strip() else 'skipRotation' if p.rotation else 'skipBlank' if not p.get_images() and not p.get_drawings() else 'pending'
                rows.append({'page':i,'status':status,'lines':0,'lowConfidence':0,'sample':''})
        return {'revision':s.revision,'pages':rows,'total':sum(x['status']=='pending' for x in rows)}
    def ocr_begin(self,sid,pages,revision):
        if self.ocr_jobs:raise EngineError('ocrBusy')
        s=self.session(sid)
        if s.revision!=revision:raise EngineError('stalePreview')
        plan=self.ocr_plan(sid,pages)
        if not plan['total']:raise EngineError('ocrNoEligible')
        job=uuid.uuid4().hex
        self.ocr_jobs[job]={'sid':sid,'revision':revision,'data':s.data,'pages':plan['pages'],'total':plan['total'],'completed':0}
        return self._ocr_status(job)
    def _ocr_job(self,job):
        task=self.ocr_jobs.get(job)
        if not task:raise EngineError('ocrJobMissing')
        if self.session(task['sid']).revision!=task['revision']:
            del self.ocr_jobs[job];raise EngineError('stalePreview')
        return task
    def _ocr_status(self,job):
        task=self._ocr_job(job)
        return {'job':job,'pages':[dict(x) for x in task['pages']],'total':task['total'],'completed':task['completed'],'done':task['completed']==task['total']}
    def ocr_step(self,job):
        task=self._ocr_job(job)
        row=next((x for x in task['pages'] if x['status']=='pending'),None)
        if not row:return self._ocr_status(job)
        try:
            with fitz.open(stream=task['data'],filetype='pdf') as d:
                p=d[row['page']];before=p.get_pixmap().samples
                result=recognize_page(p)
                if p.get_pixmap().samples!=before:raise EngineError('outsideRegionChanged')
                task['data']=d.tobytes(garbage=4,deflate=True)
            row.update(result);row['status']='recognized' if result['lines'] else 'noText'
            task['completed']+=1
            return self._ocr_status(job)
        except Exception:
            self.ocr_jobs.pop(job,None);raise
    def ocr_cancel(self,job):
        self.ocr_jobs.pop(job,None)
        self.previews={k:v for k,v in self.previews.items() if v.get('ocrJob')!=job}
        return True
    def ocr_finish(self,job):
        task=self._ocr_job(job)
        if task['completed']!=task['total']:raise EngineError('ocrIncomplete')
        try:
            changes=sum(x['lines'] for x in task['pages'])
            if not changes:raise EngineError('ocrNoText')
            s=self.session(task['sid']);diff=diff_report(s.data,task['data'],{},False)
            if diff:raise EngineError('outsideRegionChanged')
            token=uuid.uuid4().hex
            self.previews={k:v for k,v in self.previews.items() if v['sid']!=task['sid']}
            self.previews[token]={'sid':task['sid'],'revision':s.revision,'data':task['data'],'ocrJob':job,'operations':[{'type':'ocr','page':x['page'],'lines':x['lines']} for x in task['pages'] if x['lines']]}
            with fitz.open(stream=task['data'],filetype='pdf') as d:pages=len(d)
            return {'token':token,'changes':changes,'warnings':['ocrReview'],'diff':diff,'pages':pages,'revision':s.revision,'ocrPages':task['pages']}
        finally:self.ocr_jobs.pop(job,None)
    def render(self,sid,page=0,scale=1,preview=None):
        s=self.session(sid);data=s.data
        if preview:
            p=self.previews.get(preview)
            if not p or p['sid']!=sid:raise EngineError('stalePreview')
            data=p['data']
        with fitz.open(stream=data,filetype='pdf') as d:
            if not 0<=int(page)<len(d):raise EngineError('invalidPage')
            p=d[int(page)]
            return {'image':b64(png(p,max(.1,min(float(scale),2)))),'width':p.rect.width,'height':p.rect.height}
    def objects(self,sid,page=0):
        with fitz.open(stream=self.session(sid).data,filetype='pdf') as d:
            p=d[page];images=[]
            for info in p.get_image_info(xrefs=True):
                r=fitz.Rect(info['bbox'])
                images.append({'id':str(info['number']),'xref':info['xref'],'rect':list(r),'viewRect':list(r*p.rotation_matrix),'width':info['width'],'height':info['height']})
            text=spans(p)
            return {'texts':text,'images':images,'kind':'mixed' if images and text else 'native' if text else 'scan' if images else 'empty'}
    def text_backdrop(self,sid,page,id,scale=1):
        # Reuse deletion safety checks in an isolated session. Live input must not
        # mutate the document, history, or its pending commit token.
        source=self.session(sid)
        isolated=Engine()
        clone=isolated.open(b64(source.data),source.name)
        try:
            info=isolated.info(clone['id'])
            if not 0<=page<len(info['pages']):raise EngineError('invalidPage')
            if info['pages'][page]['rotation']:raise EngineError('rotatedText')
            candidate=isolated.preview(clone['id'],[{'type':'replaceText','page':page,'id':id,'text':''}])
            return isolated.render(clone['id'],page,scale,candidate['token'])
        finally:isolated.close(clone['id'])
    def text_layout(self,sid,page,id,text,size=None,bold=False,italic=False):
        with fitz.open(stream=self.session(sid).data,filetype='pdf') as doc:
            if not 0<=page<len(doc):raise EngineError('invalidPage')
            p=doc[page];span=next((x for x in spans(p) if x['id']==id),None)
            if span is None:raise EngineError('staleObject')
            return replacement_layout(p,span,text,float(size if size is not None else span['size']),font_for(text,span['font'],bold,italic),validate=False)
    def image_layer(self,sid,page,xref,scale=1):
        source=self.session(sid)
        with fitz.open(stream=source.data,filetype='pdf') as doc:
            if not 0<=page<len(doc):raise EngineError('invalidPage')
            image_placement(doc,doc[page],xref)
            raw=image_png(doc,xref)
        isolated=Engine();clone=isolated.open(b64(source.data),source.name)
        try:
            candidate=isolated.preview(clone['id'],[{'type':'deleteImage','page':page,'xref':xref}])
            backdrop=isolated.render(clone['id'],page,scale,candidate['token'])
            return {'image':b64(raw),'backdrop':backdrop['image']}
        finally:isolated.close(clone['id'])
    def search(self,sid,query):
        if not query:return []
        with fitz.open(stream=self.session(sid).data,filetype='pdf') as d:
            return [{'page':i,'rect':list(r*p.rotation_matrix)} for i,p in enumerate(d) for r in p.search_for(query)][:3000]
    def preview(self,sid,operations):
        s=self.session(sid)
        if not isinstance(operations,list) or not 0<len(operations)<=2000:raise EngineError('invalidOperation')
        kinds={o.get('type') for o in operations}
        if kinds & {'markText','deleteMarkup'} and kinds - {'markText','deleteMarkup'}:raise EngineError('invalidOperation')
        structural_kinds={'rotate','deletePage','addPage','reorder','merge','resize'}
        if kinds & structural_kinds and kinds - structural_kinds:raise EngineError('invalidOperation')
        regions={};warnings=[];structural=False;changes=0
        d=fitz.open(stream=s.data,filetype='pdf')
        try:
            for op in operations:
                kind=op.get('type'); pi=int(op.get('page',0))
                if not 0<=pi<len(d):raise EngineError('invalidPage')
                p=d[pi]
                def region(r):regions.setdefault(pi,[]).append(list(r))
                if kind=='replaceText':
                    found=next((x for x in spans(p) if x['id']==op['id']),None)
                    if not found or (op.get('expected') is not None and found['text']!=op['expected']):raise EngineError('staleObject')
                    if found['direction']!=[1.,0.]:raise EngineError('rotatedText')
                    text=str(op.get('text','')); rect=fitz.Rect(found['rect']);size=float(op.get('size',found['size']))
                    if any(ord(c)>255 for c in text) and (op.get('bold') or op.get('italic')):raise EngineError('unsupportedFontStyle')
                    font=font_for(text,found['font'],op.get('bold',False),op.get('italic',False))
                    layout=replacement_layout(p,found,text,size,font);width=layout['width'];grown=fitz.Rect(layout['rect'])
                    if replace_simple_object(d,p,found,text,op):
                        region(grown);changes+=1;continue
                    origin=fitz.Point(found['origin'])
                    # Restrict redaction to the selected glyph strip and preserve images/vectors.
                    red=fitz.Rect(rect.x0,origin.y-found['size']*.7,rect.x1,origin.y+found['size']*.05)
                    if any(x['id']!=found['id'] and fitz.Rect(x['rect']).intersects(red) for x in spans(p)):
                        raise EngineError('overlappingText')
                    if any(x.get('type')==3 and fitz.Rect(x['bbox']).intersects(rect) for x in p.get_texttrace()):
                        raise EngineError('hiddenText')
                    p.add_redact_annot(red,fill=False);p.apply_redactions(images=0,graphics=0,text=0)
                    if text:p.insert_text(origin,text,fontname=embed(p,font),fontsize=size,color=color(op.get('color',found['color'])))
                    if op.get('underline') and text:p.draw_line((rect.x0,origin.y+1),(rect.x0+width,origin.y+1),color=color(op.get('color',found['color'])),width=.7)
                    region(grown);warnings.append('fontFallback');changes+=1
                elif kind=='addText':
                    text=str(op.get('text',''));size=float(op.get('size',16));rect=fitz.Rect(op.get('rect',[60,60,p.rect.width-60,130]))
                    if not text or not 3<=size<=144 or not all(math.isfinite(v) for v in rect) or rect.is_empty or not p.rect.contains(rect):raise EngineError('invalidOperation')
                    rotation=0
                    if op.get('viewCoordinates'):
                        # Pointer placement uses displayed coordinates; PDF text APIs
                        # use unrotated page coordinates. Keep new text upright on screen.
                        rect=rect*p.derotation_matrix;rotation=p.rotation
                    spare=p.insert_textbox(rect,text,fontsize=size,fontname=embed(p,font_for(text)),color=color(op.get('color','#202124')),rotate=rotation)
                    if spare<0:raise EngineError('textBoxOverflow')
                    region(rect);changes+=1
                elif kind in ['replaceImage','deleteImage','transformImage']:
                    xref=int(op['xref']);info=unique_image(d,p,xref)
                    old=fitz.Rect(info['bbox']);region(old)
                    if op.get('expectedRect') is not None and (len(op['expectedRect'])!=4 or any(not math.isfinite(b) or abs(a-b)>.02 for a,b in zip(old,op['expectedRect']))):raise EngineError('staleObject')
                    if kind=='deleteImage':
                        # Remove the isolated drawing operation, rather than replace it
                        # with a transparent 1x1 image that remains selectable.
                        content,_,_,_=image_placement(d,p,xref);d.update_stream(content,b'')
                    elif kind=='replaceImage':p.replace_image(xref,stream=un64(op['data']))
                    else:
                        raw=image_png(d,xref)
                        im=Image.open(io.BytesIO(raw)).convert('RGBA')
                        if op.get('crop'):
                            c=op['crop']
                            if len(c)!=4 or not (0<=c[0]<c[2]<=1 and 0<=c[1]<c[3]<=1):raise EngineError('invalidOperation')
                            im=im.crop((int(c[0]*im.width),int(c[1]*im.height),int(c[2]*im.width),int(c[3]*im.height)))
                        if op.get('rotate'):im=im.rotate(int(op['rotate']),expand=True)
                        opacity=max(0,min(float(op.get('opacity',1)),1));im.putalpha(im.getchannel('A').point(lambda v:round(v*opacity)))
                        new=fitz.Rect(op.get('rect',list(old)))
                        if not all(math.isfinite(v) for v in new) or new.is_empty or not p.rect.contains(new):raise EngineError('invalidOperation')
                        if list(new)!=list(old):move_image(d,p,xref,old,new)
                        # Geometry-only edits leave the original image bytes and mask untouched.
                        if op.get('crop') or op.get('rotate') or opacity!=1:
                            buf=io.BytesIO();im.save(buf,format='PNG');p.replace_image(xref,stream=buf.getvalue())
                        region(new)
                    changes+=1
                elif kind=='insertImage':
                    rect=fitz.Rect(op.get('rect',[60,210,360,410]));p.insert_image(rect,stream=un64(op['data']),keep_proportion=True);region(rect);changes+=1
                elif kind=='rotate':p.set_rotation((p.rotation+int(op.get('degrees',90)))%360);structural=True;changes+=1
                elif kind=='deletePage':
                    if len(d)==1:raise EngineError('lastPage')
                    d.delete_page(pi);structural=True;changes+=1
                elif kind=='addPage':
                    # Candidates must remain reopenable under the same limit as open().
                    # Check every insertion, including multiple operations in one request.
                    if len(d)>=MAX_PAGES:raise EngineError('pageLimit')
                    d.new_page(pno=pi+1,width=p.rect.width,height=p.rect.height);structural=True;changes+=1
                elif kind=='reorder':
                    order=op['order']
                    if sorted(order)!=list(range(len(d))):raise EngineError('invalidOperation')
                    d.select(order);structural=True;changes+=1
                elif kind=='merge':
                    with checked(un64(op['data'])) as other:d.insert_pdf(other)
                    if len(d)>MAX_PAGES:raise EngineError('pageLimit')
                    structural=True;changes+=1
                elif kind=='resize':
                    width=float(op['width']);height=float(op['height'])
                    if not 100<=width<=3000 or not 100<=height<=3000:raise EngineError('invalidOperation')
                    # Crop/expand page box without scaling content: explicitly exposed as canvas size.
                    p.set_mediabox(fitz.Rect(0,0,width,height));structural=True;changes+=1
                elif kind in ['pageNumbers','watermark','headerFooter']:
                    for j,pg in enumerate(d):
                        text=str(j+1) if kind=='pageNumbers' else str(op.get('text',''))
                        if not text:raise EngineError('invalidOperation')
                        size=11 if kind!='watermark' else 32
                        y=pg.rect.height-26 if kind=='pageNumbers' or op.get('position')=='footer' else 30 if kind=='headerFooter' else pg.rect.height/2
                        pg.insert_text((40,y),text,fontsize=size,fontname=embed(pg,font_for(text)),color=(.5,.5,.5),fill_opacity=.25 if kind=='watermark' else 1)
                        regions.setdefault(j,[]).append([0,max(0,y-size*1.5),pg.rect.width,min(pg.rect.height,y+size)])
                    changes+=len(d)
                elif kind=='underline':
                    hits=p.search_for(str(op['query']))
                    for r in hits:
                        a=p.add_underline_annot(r);a.set_colors(stroke=color(op.get('color','#d6334c')));a.update();region(r)
                    changes+=len(hits)
                elif kind=='markText':
                    if op.get('revision')!=s.revision:raise EngineError('stalePreview')
                    style=op.get('style');hits,_=markup_hits(p,op.get('query'),style)
                    tint=color(op.get('color','#ffe066'));opacity=op.get('opacity',.45)
                    if type(opacity) not in (int,float) or not math.isfinite(opacity) or not .1<=opacity<=1:raise EngineError('invalidMarkup')
                    if changes+len(hits)>2000:raise EngineError('markupLimit')
                    for quad in hits:
                        a={'highlight':p.add_highlight_annot,'underline':p.add_underline_annot,'strikeout':p.add_strikeout_annot}[style](quad)
                        a.set_colors(stroke=tint);a.set_info(title='OpenPDF Studio',content=op['query'],subject='OpenPDF Studio / keyword markup')
                        d.xref_set_key(a.xref,'OpenPDFStudio','('+MARKUP_TAG+')');a.set_flags(fitz.PDF_ANNOT_IS_PRINT);a.update(opacity=opacity)
                        region(a.rect);changes+=1
                elif kind=='deleteMarkup':
                    if op.get('revision')!=s.revision:raise EngineError('stalePreview')
                    xref=op.get('xref')
                    if type(xref) is not int or xref not in [a.xref for a in p.annots() or []]:raise EngineError('staleObject')
                    a=p.load_annot(xref)
                    if d.xref_get_key(xref,'OpenPDFStudio')!=('string',MARKUP_TAG) or a.info.get('id','')!=op.get('id') or a.type[0] not in MARKUP_TYPES.values():raise EngineError('staleObject')
                    region(a.rect);p.delete_annot(a);changes+=1
                elif kind=='ocr':
                    result=recognize_page(p);warnings.append('ocrReview');changes+=result['lines']
                else:raise EngineError('invalidOperation')
            if not changes:raise EngineError('noMatches')
            after=d.tobytes(garbage=4,deflate=True)
        finally:d.close()
        markup_only=kinds<={'markText','deleteMarkup'}
        if markup_only:
            with fitz.open(stream=s.data,filetype='pdf') as original,fitz.open(stream=after,filetype='pdf') as modified:
                for i in range(len(original)):
                    if [original.xref_stream(x) for x in original[i].get_contents()]!=[modified.xref_stream(x) for x in modified[i].get_contents()] or original[i].get_pixmap(alpha=True,annots=False).samples!=modified[i].get_pixmap(alpha=True,annots=False).samples:raise EngineError('outsideRegionChanged')
        diff=diff_report(s.data,after,regions,structural,alpha=markup_only)
        if any(x.get('outsideRegionChanged') for x in diff):raise EngineError('outsideRegionChanged',diff)
        token=uuid.uuid4().hex
        # Only latest candidate per document is retained.
        self.previews={k:v for k,v in self.previews.items() if v['sid']!=sid}
        self.previews[token]={'sid':sid,'revision':s.revision,'data':after,'operations':[{k:v for k,v in o.items() if k!='data'} for o in operations]}
        return {'token':token,'changes':changes,'warnings':sorted(set(warnings)),'markupValidated':markup_only,'diff':diff,'pages':len(fitz.open(stream=after,filetype='pdf')),'revision':s.revision}
    def commit(self,sid,token):
        s=self.session(sid);p=self.previews.get(token)
        if not p or p['sid']!=sid or p['revision']!=s.revision:raise EngineError('stalePreview')
        s.undo.append(s.data);s.undo=s.undo[-15:]
        while len(s.undo)>1 and sum(map(len,s.undo))>128*1024*1024:s.undo.pop(0)
        s.data=p['data'];s.redo=[];s.revision+=1;s.log.append({'revision':s.revision,'operations':p['operations']})
        del self.previews[token];return self.info(sid)
    def history(self,sid,direction):
        s=self.session(sid)
        if direction not in ['undo','redo']:raise EngineError('invalidOperation')
        source,target=(s.undo,s.redo) if direction=='undo' else (s.redo,s.undo)
        if not source:raise EngineError('noHistory')
        target.append(s.data);s.data=source.pop();s.revision+=1;return self.info(sid)
    def plan(self,sid,instruction):
        # Exact bounded grammar. This planner has no code execution or network provider.
        t=instruction.strip();operations=[]
        m=re.fullmatch(r'(?:rotate page|旋转第?)\s*(\d+)\s*(?:页)?(?:\s*(?:by|顺时针)\s*(90|180|270)\s*(?:degrees|度)?)?',t,re.I)
        if m:operations=[{'type':'rotate','page':int(m[1])-1,'degrees':int(m[2] or 90)}]
        m2=re.fullmatch(r'(?:underline|下划线)\s+["“](.+?)["”]',t,re.I)
        if m2:
            operations=[{'type':'underline','page':i,'query':m2[1]} for i in range(len(self.info(sid)['pages']))]
        m3=re.fullmatch(r'(?:replace|替换)\s+["“](.+?)["”]\s*(?:with|为)\s*["“](.*?)["”]',t,re.I)
        if m3:operations=self.replace_plan(sid,m3[1],m3[2])
        if not operations:raise EngineError('unsupportedInstruction')
        return {'operations':operations,'provider':'local-rules'}
    def replace_plan(self,sid,query,replacement):
        if not query:raise EngineError('noMatches')
        ops=[]
        with fitz.open(stream=self.session(sid).data,filetype='pdf') as d:
            for i,p in enumerate(d):
                # Reverse reading order minimizes span index changes; expected text still guards stale selectors.
                for s in reversed(spans(p)):
                    if query in s['text']:ops.append({'type':'replaceText','page':i,'id':s['id'],'expected':s['text'],'text':s['text'].replace(query,replacement)})
        if not ops:raise EngineError('noMatches')
        return ops
    def markup_plan(self,sid,query,pages,style='highlight'):
        validate_markup(query,style);s=self.session(sid);rows=[];operations=[];total=0;duplicates=0
        with fitz.open(stream=s.data,filetype='pdf') as d:
            if not isinstance(pages,list) or not pages or any(type(i) is not int or not 0<=i<len(d) for i in pages):raise EngineError('invalidPage')
            for i in sorted(set(pages)):
                hits,existing=markup_hits(d[i],query,style);total+=len(hits);duplicates+=existing
                if total>2000:raise EngineError('markupLimit')
                rows.append({'page':i,'count':len(hits),'duplicates':existing})
                if hits:operations.append({'type':'markText','page':i,'query':query,'style':style,'revision':s.revision})
        return {'revision':s.revision,'pages':rows,'total':total,'duplicates':duplicates,'operations':operations}
    def markups(self,sid):
        s=self.session(sid);rows=[]
        with fitz.open(stream=s.data,filetype='pdf') as d:
            for i,p in enumerate(d):
                for a in p.annots() or []:
                    if d.xref_get_key(a.xref,'OpenPDFStudio')!=('string',MARKUP_TAG):continue
                    style=next((k for k,v in MARKUP_TYPES.items() if a.type[0]==v),None)
                    if style:rows.append({'page':i,'xref':a.xref,'id':a.info.get('id',''),'text':a.info.get('content',''),'style':style,'rect':list(a.rect)})
        return {'revision':s.revision,'items':rows}
    def conversion_info(self,sid):
        with fitz.open(stream=self.session(sid).data,filetype='pdf') as d:
            image_only=sum(not p.get_text().strip() and bool(p.get_images() or p.get_drawings()) for p in d)
            return {'pages':len(d),'imageOnlyPages':image_only}
    def export(self,sid,format='pdf',pages=None):
        s=self.session(sid)
        d=fitz.open(stream=s.data,filetype='pdf')
        try:
            if pages is not None:
                if not pages or any(not isinstance(x,int) or x<0 or x>=len(d) for x in pages):raise EngineError('invalidPage')
                d.select(pages)
            out=io.BytesIO();ext=format
            if format=='pdf':data=d.tobytes(garbage=4,deflate=True)
            elif format=='txt':data='\n\n'.join(p.get_text(sort=True) for p in d).encode()
            elif format in ['png','split']:
                ext='zip'
                with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
                    for i,p in enumerate(d):
                        if format=='png':z.writestr(f'page-{i+1}.png',png(p,2))
                        else:
                            part=fitz.open();part.insert_pdf(d,from_page=i,to_page=i);z.writestr(f'page-{i+1}.pdf',part.tobytes());part.close()
                data=out.getvalue()
            elif format=='docx':
                from docx import Document
                doc=Document()
                for i,p in enumerate(d):
                    if i:doc.add_page_break()
                    for line in p.get_text(sort=True).splitlines():doc.add_paragraph(line)
                doc.save(out);data=out.getvalue()
            elif format=='xlsx':
                from openpyxl import Workbook
                from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
                w=Workbook();w.remove(w.active)
                for i,p in enumerate(d):
                    sheet=w.create_sheet(f'Page {i+1}');tables=p.find_tables()
                    rows=[row for table in tables.tables for row in table.extract()] or [[line] for line in p.get_text(sort=True).splitlines()]
                    for row in rows:
                        sheet.append([ILLEGAL_CHARACTERS_RE.sub('',str(x or '')) for x in row])
                        for cell in sheet[sheet.max_row]:cell.data_type='s'
                w.save(out);data=out.getvalue()
            elif format=='pptx':
                from pptx import Presentation
                from pptx.util import Inches
                r=Presentation();r.slide_width=Inches(d[0].rect.width/72);r.slide_height=Inches(d[0].rect.height/72)
                for p in d:r.slides.add_slide(r.slide_layouts[6]).shapes.add_picture(io.BytesIO(png(p,2)),0,0,width=r.slide_width,height=r.slide_height)
                r.save(out);data=out.getvalue()
            elif format=='log':ext='json';data=json.dumps(s.log,ensure_ascii=False,indent=2).encode()
            else:raise EngineError('invalidOperation')
            return {'data':b64(data),'name':Path(s.name).stem+'-edited.'+ext,'extension':ext}
        finally:d.close()
    def import_file(self,data,name):
        raw=un64(data);ext=Path(name).suffix.lower()
        if ext=='.pdf':return self.open(data,name)
        if ext=='.docx':
            from docx import Document
            import html
            doc=Document(io.BytesIO(raw));body=''.join('<p>'+html.escape(p.text)+'</p>' for p in doc.paragraphs)
            for table in doc.tables:body+='<table>'+''.join('<tr>'+''.join('<td>'+html.escape(c.text)+'</td>' for c in row.cells)+'</tr>' for row in table.rows)+'</table>'
            out=io.BytesIO();writer=fitz.DocumentWriter(out);story=fitz.Story(html=body,user_css='body{font-family:sans-serif;font-size:12pt} td{border:1px solid gray}')
            more=True
            while more:
                device=writer.begin_page(fitz.paper_rect('a4'));more,_=story.place(fitz.Rect(40,40,555,802));story.draw(device);writer.end_page()
            writer.close();pdf=out.getvalue()
        else:
            try:
                with Image.open(io.BytesIO(raw)) as im:im.verify()
                img=fitz.open(stream=raw);pdf=img.convert_to_pdf();img.close()
            except Exception:raise EngineError('invalidImage')
        return self.open(b64(pdf),Path(name).stem+'.pdf')
